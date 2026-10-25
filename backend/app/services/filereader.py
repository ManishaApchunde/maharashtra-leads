"""Read an uploaded CSV or Excel file into a list of dicts with normalised column names."""
import csv
import io
import re

BUSINESS_COLUMNS = [
    "shop_name", "contact_name", "phone", "alternate_phone", "whatsapp_number", "email", "address",
    "area", "village", "city", "taluka", "district", "pincode", "latitude", "longitude",
    "google_maps_url", "category", "source", "source_url",
]
OPTIONAL_COLUMNS = ["website", "source_place_id", "rating", "review_count", "keyword",
                    "discovered_keyword", "whatsapp_same_as_phone"]
REQUIRED_COLUMNS = ["shop_name"]

# Common header variants -> our names
HEADER_ALIASES = {
    "name": "shop_name", "business_name": "shop_name", "shop": "shop_name", "shopname": "shop_name",
    "owner": "contact_name", "owner_name": "contact_name", "contact": "contact_name", "contact_person": "contact_name",
    "mobile": "phone", "mobile_number": "phone", "phone_number": "phone", "contact_number": "phone",
    "alt_phone": "alternate_phone", "alternate_mobile": "alternate_phone", "phone2": "alternate_phone",
    "whatsapp": "whatsapp_number", "whatsapp_no": "whatsapp_number",
    "pin": "pincode", "pin_code": "pincode", "postal_code": "pincode", "zip": "pincode",
    "lat": "latitude", "lng": "longitude", "lon": "longitude", "long": "longitude",
    "tehsil": "taluka", "tahsil": "taluka", "tal": "taluka",
    "town": "city", "locality": "area", "map_url": "google_maps_url", "maps_url": "google_maps_url",
    "place_id": "source_place_id", "reviews": "review_count",
}


def norm_header(h: str, aliases: bool = True) -> str:
    h = re.sub(r"[^a-z0-9]+", "_", (h or "").strip().lower()).strip("_")
    return HEADER_ALIASES.get(h, h) if aliases else h


def read_table(content: bytes, filename: str, aliases: bool = True) -> tuple[list[str], list[dict]]:
    """aliases=True maps business-file headers like 'mobile' -> 'phone'."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = [norm_header(str(h or ""), aliases) for h in next(rows, [])]
        data = []
        for r in rows:
            if r is None or all(v is None or str(v).strip() == "" for v in r):
                continue
            data.append({headers[i]: ("" if v is None else str(v)) for i, v in enumerate(r) if i < len(headers)})
        return headers, data

    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    reader = csv.DictReader(io.StringIO(text))
    headers = [norm_header(h, aliases) for h in (reader.fieldnames or [])]
    reader.fieldnames = headers
    data = [r for r in reader if any((v or "").strip() for v in r.values() if isinstance(v, str))]
    return headers, data


def validate_headers(headers: list[str], required: list[str]) -> tuple[list[str], list[str]]:
    """Returns (missing_required, ignored_unknown)."""
    known = set(BUSINESS_COLUMNS) | set(OPTIONAL_COLUMNS)
    missing = [c for c in required if c not in headers]
    ignored = [h for h in headers if h and h not in known]
    return missing, ignored
